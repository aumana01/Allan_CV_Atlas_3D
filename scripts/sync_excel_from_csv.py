#!/usr/bin/env python3
"""Synchronize master Excel with reviewed CSV fields without rebuilding the workbook."""
from pathlib import Path
import csv
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CSV_FILE = ROOT / "data" / "proyectos.csv"
EXCEL_FILE = ROOT / "data" / "proyectos.xlsx"
FIELDS = ("descripcion", "participacion_detalle", "rol", "categoria", "tecnologias")
IDS_TO_SYNC = {"18", "19"}

with CSV_FILE.open(encoding="utf-8-sig", newline="") as f:
    records = {str(row["id"]).strip(): row for row in csv.DictReader(f)}

missing = IDS_TO_SYNC - records.keys()
if missing:
    raise SystemExit(f"Missing source project ID(s): {sorted(missing)}")

wb = load_workbook(EXCEL_FILE)
ws = wb["Proyectos"] if "Proyectos" in wb.sheetnames else wb.active
header = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value is not None}
for field in ("id",) + FIELDS:
    if field not in header:
        raise SystemExit(f"Missing Excel column: {field}")

updated = set()
for row_idx in range(2, ws.max_row + 1):
    id_value = ws.cell(row_idx, header["id"]).value
    project_id = str(id_value).strip() if id_value is not None else ""
    if project_id not in IDS_TO_SYNC:
        continue
    for field in FIELDS:
        ws.cell(row_idx, header[field]).value = records[project_id][field]
    updated.add(project_id)

if updated != IDS_TO_SYNC:
    raise SystemExit(f"Could not find expected Excel IDs: {sorted(IDS_TO_SYNC - updated)}")
wb.save(EXCEL_FILE)
print("Excel master synchronized:", ", ".join(sorted(updated)))
